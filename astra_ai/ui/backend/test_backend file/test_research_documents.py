"""Real document imports, immutable originals and revision-checked PDF previews."""
import asyncio,io,hashlib
import pytest
from PIL import Image
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader
from astra_ai.core.research_learning import ResearchArtifacts,ResearchSearch
from astra_ai.core.research_documents import DocumentFiles,create_entry,compose_documentary
CTX={'profile_id':'owner','project_id':'project','session_id':'main'}

def fixture_pdf(scan=False):
    stream=io.BytesIO();canvas=Canvas(stream,pagesize=(600,800))
    canvas.setFont('Helvetica',14)
    if not scan:canvas.drawString(50,740,'Research document heading')
    canvas.drawImage(ImageReader(Image.new('RGB',(200,100),'teal')),50,500,width=200,height=100)
    canvas.rect(400,50,100,100,stroke=1,fill=0);canvas.save();return stream.getvalue()

def setup(tmp_path,body=None,name='sample.pdf'):
    artifacts=ResearchArtifacts(tmp_path/'reports.json')
    entry=create_entry(artifacts,CTX,'documentary','Explain this uploaded research','document')
    files=DocumentFiles(artifacts);identifier=entry['research_id'];row=files.upload(identifier,CTX,name,body or fixture_pdf())
    files.process(identifier,CTX,row['id'])
    return artifacts,files,identifier,files.get(identifier,CTX,row['id'])

def test_pdf_import_real_text_image_pages_and_scoped_original(tmp_path):
    store,files,identifier,row=setup(tmp_path)
    page=files.page(identifier,CTX,row['id'])
    assert page['page_count']==1 and 'Research document heading' in page['text']
    assert {o['type'] for o in page['objects']}=={'text','image'}
    assert row['images'] and not row['needs_ocr']
    assert store.document(identifier,CTX,bounded=True)['attachment_images'][0]['source_id']=='S1'
    assert hashlib.sha256(files.original(identifier,CTX,row['id'])[0]).hexdigest()==row['revision']
    with pytest.raises(ValueError):files.page(identifier,{**CTX,'project_id':'other'},row['id'])
    with pytest.raises(ValueError):files.folder('../escape',row['id'])

def test_edit_preview_apply_undo_and_unaffected_region(tmp_path):
    store,files,identifier,row=setup(tmp_path);original=files.original(identifier,CTX,row['id'])[0]
    image=next(o for o in files.page(identifier,CTX,row['id'])['objects'] if o['type']=='image')
    result=files.edit_pdf(identifier,CTX,row['id'],row['revision'],[{'action':'move_object','page':1,'object_id':image['id'],'rect':[70,450,200,100]}],'move')
    assert result['status']=='preview' and not files.get(identifier,CTX,row['id']).get('pdf_revision')
    edited=files.original(identifier,CTX,row['id'],result['revision'])[0]
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(original) as doc:before=doc[0].render(scale=1).to_pil().copy()
    with pdfium.PdfDocument(edited) as doc:after=doc[0].render(scale=1).to_pil().copy()
    assert before.crop((0,0,600,120)).tobytes()==after.crop((0,0,600,120)).tobytes()
    assert before.crop((390,630,520,760)).tobytes()==after.crop((390,630,520,760)).tobytes()
    assert before.tobytes()!=after.tobytes()
    before.save(tmp_path/'original.png');after.save(tmp_path/'edited.png')
    files.edit_pdf(identifier,CTX,row['id'],row['revision'],[],'move',True)
    assert files.get(identifier,CTX,row['id'])['pdf_revision']==result['revision']
    with pytest.raises(ValueError,match='revision_changed'):files.edit_pdf(identifier,CTX,row['id'],row['revision'],[{'action':'add_page','page':1}],'stale')
    files.restore_pdf(identifier,CTX,row['id'],result['revision'],row['revision'])
    assert files.original(identifier,CTX,row['id'])[0]==original

def test_text_replacement_overflow_and_page_insertion(tmp_path):
    _,files,identifier,row=setup(tmp_path)
    text=next(o for o in files.page(identifier,CTX,row['id'])['objects'] if o['type']=='text')
    with pytest.raises(ValueError,match='overflow'):files.edit_pdf(identifier,CTX,row['id'],row['revision'],[{'action':'replace_text','page':1,'object_id':text['id'],'text':'Too long '*100}],'overflow')
    result=files.edit_pdf(identifier,CTX,row['id'],row['revision'],[{'action':'replace_text','page':1,'object_id':text['id'],'text':'Updated heading'},{'action':'add_page','page':1},{'action':'insert_text','page':2,'rect':[50,600,450,150],'text':'New supported text in an added page.'}],'text-edit')
    page=files.page(identifier,CTX,row['id'],1,result['revision'])
    assert 'Updated heading' in page['text'] and page['page_count']==2
    assert 'New supported text' in files.page(identifier,CTX,row['id'],2,result['revision'])['text']

def test_scans_encryption_malformed_and_file_limits(tmp_path):
    _,files,identifier,row=setup(tmp_path,fixture_pdf(True))
    assert row['needs_ocr'] and not any(o['type']=='text' for o in files.page(identifier,CTX,row['id'])['objects'])
    broken=files.upload(identifier,CTX,'broken.pdf',b'%PDF-corrupt')
    with pytest.raises(ValueError):files.process(identifier,CTX,broken['id'])
    assert files.get(identifier,CTX,broken['id'])['status']=='failed'
    from pypdf import PdfReader,PdfWriter
    writer=PdfWriter();writer.append(PdfReader(io.BytesIO(fixture_pdf())));writer.encrypt('locked');out=io.BytesIO();writer.write(out)
    locked=files.upload(identifier,CTX,'encrypted.pdf',out.getvalue())
    with pytest.raises(ValueError):files.process(identifier,CTX,locked['id'])
    with pytest.raises(ValueError):files.upload(identifier,CTX,'unsupported.exe',b'unsupported')
    assert files.upload(identifier,CTX,'same.pdf',fixture_pdf(True))['id'] # distinct generation metadata allowed
    with pytest.raises(ValueError):files.upload(identifier,CTX,'huge.txt',b'x'*(25*1024*1024+1))

def test_docx_text_tables_images_and_bounded_blocks(tmp_path):
    from docx import Document
    doc=Document();doc.add_heading('Evidence heading',level=1);doc.add_paragraph('Supported body. '*1500)
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Topic';table.cell(0,1).text='Finding'
    image=io.BytesIO();Image.new('RGB',(160,100),'orange').save(image,format='PNG');image.seek(0);doc.add_picture(image)
    body=io.BytesIO();doc.save(body)
    _,files,identifier,row=setup(tmp_path,body.getvalue(),'document.docx')
    page=files.page(identifier,CTX,row['id'])
    assert len(page['text'])<=8000 and page['next_offset'] and row['images']
    rest=files.page(identifier,CTX,row['id'],offset=page['next_offset'])
    assert rest['offset']==8000

def test_documentary_reuses_sources_and_quota_checkpoint(tmp_path):
    store,files,identifier,row=setup(tmp_path)
    class Model:
        async def report(self,brief,sources,kind):
            assert kind=='documentary' and sources[0]['attachment_id']==row['id']
            return '# A concise documentary title\n\n## Introduction\n\nSupported explanation [S1].\n\n## Conclusion\n\nA bounded conclusion [S1].'
    search=ResearchSearch(store,lambda k:None)
    result=asyncio.run(compose_documentary(store,search,Model(),CTX,identifier))
    assert result['status']=='completed' and result['kind']=='documentary'
    assert store.get(identifier,CTX)['artifact']['report'].startswith('# A concise')
    assert files.original(identifier,CTX,row['id'])[0].startswith(b'%PDF-')
    second=asyncio.run(compose_documentary(store,search,None,CTX,identifier))
    assert second['research_revision']==result['research_revision']

def test_document_api_scopes_uploads_and_queues_only_handles(tmp_path):
    from fastapi import FastAPI,HTTPException
    from fastapi.testclient import TestClient
    from astra_ai.ui.backend.document_api import install,run_document_job
    store=ResearchArtifacts(tmp_path/'reports.json');entry=create_entry(store,CTX,'documentary','Document fixture','op')
    calls=[]
    class Jobs:
        def submit(self,actions,context,operation_id,goal):calls.append((actions,context,operation_id));return {'status':'queued','job_id':'job'}
    def context(request):
        if request.headers.get('X-Nova-Client')!='test-owner':raise HTTPException(403,'No owner connection')
        return CTX
    app=FastAPI();search=ResearchSearch(store,lambda k:None);install(app,lambda:(store,search),context,lambda:Jobs());client=TestClient(app)
    path='/api/research/'+entry['research_id'];headers={'X-Nova-Client':'test-owner','X-File-Name':'evidence.txt'}
    assert client.post(path+'/attachments',content=b'Supported source text.').status_code==403
    result=client.post(path+'/attachments',content=b'Supported source text.',headers=headers).json();file_id=result['attachment']['id']
    args=calls[0][0][0]['arguments'];assert set(args)=={'action','research_id','attachment_id'}
    done=asyncio.run(run_document_job(store,search,None,asyncio.Semaphore(1),args,CTX));assert done['status']=='completed'
    assert client.get(path+'/attachments/'+file_id+'/page',headers=headers).json()['text']=='Supported source text.'
    assert client.get(path+'/attachments/../original',headers=headers).status_code!=200
    assert client.get('/api/research/'+'a'*32+'/attachments',headers=headers).status_code==404

def test_authored_preview_is_sanitized_and_cannot_overwrite_changed_draft(tmp_path):
    from astra_ai.ui.backend.document_api import authored_preview,authored_apply
    store,files,identifier,row=setup(tmp_path)
    preview=authored_preview(files,identifier,CTX,'# Edited copy\n\n<script>injected()</script>Supported text.','', 'preview')
    assert '<script>' not in preview['html']
    result=authored_apply(files,identifier,CTX,'preview');assert result['status']=='completed'
    assert store.get(identifier,CTX)['artifact']['report']=='' # original is unchanged
    next_preview=authored_preview(files,identifier,CTX,'# Another proposed change',result['draft']['revision'],'next')
    store.save_draft(identifier,CTX,store.get(identifier,CTX)['artifact']['revision'],result['draft']['revision'],'<p>Owner edit</p>')
    assert authored_apply(files,identifier,CTX,'next')['status']=='conflict'

def test_documentary_pending_releases_budget_and_reuses_upload_after_resume(tmp_path):
    from astra_ai.core.background_budget import BudgetPaused
    store,files,identifier,row=setup(tmp_path)
    class Paused:
        async def report(self,*args,**kwargs):raise BudgetPaused(9999999999,'daily_budget')
    search=ResearchSearch(store,lambda k:None)
    first=asyncio.run(compose_documentary(store,search,Paused(),CTX,identifier))
    assert first['status']=='waiting_for_quota' and store.request_state(identifier,CTX)['phase']=='waiting_for_quota'
    assert files.get(identifier,CTX,row['id'])['status']=='ready'
    class Ready:
        async def report(self,brief,sources,kind):return '# Documentary\n\nA supported document [S1].'
    resumed=asyncio.run(compose_documentary(store,ResearchSearch(store,lambda k:None),Ready(),CTX,identifier))
    assert resumed['status']=='completed' and len(DocumentFiles(store).listing(identifier,CTX))==1
